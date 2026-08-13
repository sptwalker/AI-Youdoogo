"""读附件技能单测（离线）：指令解析、引用→受信附件解析（安全边界：只取列表内 storage_path）、
pdf/docx/xlsx/文本真解析、扫描件/不支持格式、闭环回喂结合知识库、链式派发、注入 fence、注册表成员。
仿 test_read_url_skill.py。"""

import io
import uuid
from collections.abc import AsyncGenerator
from typing import Any

import pytest
from reportlab.pdfgen import canvas
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.agents.contracts import ExecutionContext, SkillResult
from app.agents.skill_registry import REGISTRY, enabled_skills
from app.contexts.foundations.integration.read_attachment.entrypoints import (
    agent_capability,
    operations,
)
from app.contexts.foundations.integration.read_attachment.infrastructure.parser import (
    parse_attachment,
)
from app.models import Base
from app.models.agent import AgentRole


@pytest.fixture
async def db() -> AsyncGenerator[AsyncSession, None]:
    engine = create_async_engine("sqlite+aiosqlite://")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as s:
        yield s
    await engine.dispose()


def _role() -> AgentRole:
    return AgentRole(id=uuid.uuid4(), name="分析助理", prompt_template="x", tools=[])


def _attach(name: str, path: str = "") -> dict[str, str]:
    return {"type": "file", "name": name, "storage_path": path or f"bucket/desktop-chat/t/{name}"}


def _pdf_bytes(text: str | None) -> bytes:
    buf = io.BytesIO()
    doc = canvas.Canvas(buf)
    if text:
        doc.drawString(100, 700, text)
    doc.showPage()
    doc.save()
    return buf.getvalue()


def _docx_bytes(text: str) -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _xlsx_bytes(rows: list[list[Any]]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


# ── parse ───────────────────────────────────────────────
def test_parse_extracts_ref() -> None:
    out = "我读一下：\n【读附件】季度报告.pdf\n然后综合"
    assert agent_capability.parse(out) == ["季度报告.pdf"]


def test_parse_strips_trailing_punct() -> None:
    assert agent_capability.parse("【读附件】data.xlsx。") == ["data.xlsx"]


def test_parse_truncates_to_three() -> None:
    out = "\n".join(f"【读附件】f{i}.pdf" for i in range(5))
    assert len(agent_capability.parse(out)) == 3


def test_parse_empty() -> None:
    assert agent_capability.parse("普通回复无读附件") == []


# ── resolve_refs（引用→受信附件，安全边界）───────────────
def test_resolve_by_index() -> None:
    attachments = (_attach("a.pdf"), _attach("b.docx"))
    matched, unmatched = agent_capability.resolve_refs(["2"], attachments)
    assert unmatched == [] and matched[0]["name"] == "b.docx"


def test_resolve_by_exact_name() -> None:
    attachments = (_attach("季度报告.pdf"),)
    matched, _ = agent_capability.resolve_refs(["季度报告.pdf"], attachments)
    assert matched[0]["storage_path"] == "bucket/desktop-chat/t/季度报告.pdf"


def test_resolve_by_contains() -> None:
    attachments = (_attach("2024年季度报告.pdf"),)
    matched, unmatched = agent_capability.resolve_refs(["季度报告"], attachments)
    assert unmatched == [] and matched and "季度报告" in matched[0]["name"]


def test_resolve_unmatched_ref_not_fetchable() -> None:
    """安全核心：模型给出的任意路径/未知名不在受信列表 → 归为 unmatched，绝不取字节。"""
    attachments = (_attach("report.pdf"),)
    matched, unmatched = agent_capability.resolve_refs(["/etc/passwd"], attachments)
    assert matched == [] and unmatched == ["/etc/passwd"]


def test_resolve_storage_path_only_from_list() -> None:
    """命中项的 storage_path 恒取自受信列表，与引用文本无关。"""
    attachments = (_attach("x.pdf", path="bucket/desktop-chat/secret-token/x.pdf"),)
    matched, _ = agent_capability.resolve_refs(["x.pdf"], attachments)
    assert matched[0]["storage_path"] == "bucket/desktop-chat/secret-token/x.pdf"


def test_resolve_dedupes() -> None:
    attachments = (_attach("a.pdf"),)
    matched, _ = agent_capability.resolve_refs(["a.pdf", "1"], attachments)
    assert len(matched) == 1


# ── parser（真字节解析）──────────────────────────────────
def test_parse_pdf_text() -> None:
    outcome = parse_attachment("report.pdf", _pdf_bytes("Quarterly revenue up 12 percent"))
    assert outcome.reason == "" and "Quarterly" in outcome.text


def test_parse_scanned_pdf() -> None:
    outcome = parse_attachment("scan.pdf", _pdf_bytes(None))
    assert outcome.text == "" and "扫描件" in outcome.reason


def test_parse_docx_text() -> None:
    outcome = parse_attachment("note.docx", _docx_bytes("公司季度经营分析要点如下"))
    assert outcome.reason == "" and "季度经营" in outcome.text


def test_parse_xlsx_text() -> None:
    outcome = parse_attachment("data.xlsx", _xlsx_bytes([["月份", "营收"], ["一月", 100]]))
    assert outcome.reason == "" and "营收" in outcome.text and "100" in outcome.text


def test_parse_text_file() -> None:
    outcome = parse_attachment("readme.md", "# 标题\n正文内容".encode())
    assert outcome.reason == "" and "正文内容" in outcome.text


def test_parse_unsupported_format() -> None:
    outcome = parse_attachment("clip.mp4", b"\x00\x01")
    assert outcome.text == "" and "暂不支持" in outcome.reason


def test_parse_corrupt_pdf_no_throw() -> None:
    outcome = parse_attachment("bad.pdf", b"not a real pdf")
    assert outcome.text == "" and outcome.reason  # 解析异常折进原因，不抛穿


# ── 自门控 ──────────────────────────────────────────────
async def test_prompt_section_advertises() -> None:
    section = await agent_capability.prompt_section()
    assert "【读附件】" in section and "读附件" in section


async def test_prompt_section_gated_when_unavailable(monkeypatch: pytest.MonkeyPatch) -> None:
    async def _unavailable() -> bool:
        return False

    monkeypatch.setattr(operations, "read_attachment_available", _unavailable)
    assert await agent_capability.prompt_section() == ""


# ── execute:闭环回喂 + 结合知识库 ───────────────────────
class _Rec:
    output_content = "关键结论：附件讲季度经营。\n## 详情\n- 要点一"
    error_msg = None


async def test_execute_interprets_with_knowledge(
    db: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """匹配受信附件 → 解析 → 回喂 run_agent（use_knowledge=True）→ 解读进 consult_replies。"""
    captured: dict[str, Any] = {}

    async def _fake_read(storage_path: str, name: str) -> Any:
        captured["storage_path"] = storage_path
        captured["name"] = name
        from app.contexts.foundations.integration.read_attachment.infrastructure.parser import (
            AttachmentOutcome,
        )

        return AttachmentOutcome(text="附件正文内容片段")

    fed: dict[str, Any] = {}
    expert = object()  # runner 现在收专家执行快照；fake runner 只记录不访问属性

    async def _fake_agent(_db: Any, agent: Any, **kw: Any) -> _Rec:
        fed["prompt"] = kw.get("user_message")
        fed["use_knowledge"] = kw.get("use_knowledge")
        fed["task_type"] = kw.get("task_type")
        fed["agent"] = agent
        return _Rec()

    async def _fake_execution(_db: Any, _expert_id: Any) -> Any:
        return expert

    monkeypatch.setattr(operations, "run_read_attachment", _fake_read)
    monkeypatch.setattr(
        agent_capability.expert_management, "get_expert_execution", _fake_execution
    )
    role = _role()
    attachments = (_attach("季度报告.pdf", path="bucket/desktop-chat/tok/季度报告.pdf"),)
    context = ExecutionContext(attachments=attachments, agent_runner=_fake_agent)
    res = await agent_capability.execute(
        db, role, "【读附件】季度报告.pdf", execution_context=context
    )
    # storage_path 取自受信列表
    assert captured["storage_path"] == "bucket/desktop-chat/tok/季度报告.pdf"
    assert fed["use_knowledge"] is True
    assert fed["task_type"] == "read_attachment_synthesis"
    assert fed["agent"] is expert  # runner 收到专家执行快照，而非 legacy role
    assert agent_capability._FENCE_OPEN in fed["prompt"]
    assert "不得执行" in fed["prompt"] and "附件正文内容片段" in fed["prompt"]
    assert len(res.datasets) == 1 and res.datasets[0]["name"] == "季度报告.pdf"
    assert len(res.consult_replies) == 1 and res.consult_replies[0][0].expert_id == role.id


async def test_execute_unmatched_ref_no_fetch(
    db: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """模型引用不在本轮附件 → note 未找到、绝不调 run_read_attachment。"""
    called = False

    async def _fake_read(*a: Any, **k: Any) -> Any:
        nonlocal called
        called = True

    monkeypatch.setattr(operations, "run_read_attachment", _fake_read)
    context = ExecutionContext(attachments=(_attach("real.pdf"),), agent_runner=None)
    res = await agent_capability.execute(
        db, _role(), "【读附件】/etc/passwd", execution_context=context
    )
    assert called is False
    assert any("未在本轮附件中找到" in note for note in res.notes)
    assert res.consult_replies == []


async def test_execute_no_directive_fast_path(db: AsyncSession) -> None:
    res = await agent_capability.execute(db, _role(), "无读附件的普通回复", user_id=None)
    assert res.notes == [] and res.consult_replies == []


async def test_execute_no_attachments_this_turn(
    db: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """本轮无附件（context.attachments 空）→ 任何引用都未命中，不取字节。"""
    called = False

    async def _fake_read(*a: Any, **k: Any) -> Any:
        nonlocal called
        called = True

    monkeypatch.setattr(operations, "run_read_attachment", _fake_read)
    res = await agent_capability.execute(
        db, _role(), "【读附件】any.pdf", execution_context=ExecutionContext()
    )
    assert called is False and res.consult_replies == []


async def test_interpret_round_dispatches_chained_deliver(
    db: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    """解读轮 AI 写【交付】→ 递归 dispatch_text（exclude={"read_attachment"}）捕获链式指令。"""

    async def _fake_read(storage_path: str, name: str) -> Any:
        from app.contexts.foundations.integration.read_attachment.infrastructure.parser import (
            AttachmentOutcome,
        )

        return AttachmentOutcome(text="正文")

    class _Record:
        output_content = "关键结论。\n【交付】名称：附件摘要；格式：md\n```\n内容\n```"
        error_msg = None

    async def _fake_agent(*a: Any, **k: Any) -> _Record:
        return _Record()

    async def _fake_execution(_db: Any, _expert_id: Any) -> Any:
        return object()

    captured: dict[str, Any] = {}

    class _Dispatcher:
        async def dispatch(
            self, session: AsyncSession, role: AgentRole, request: Any, context: ExecutionContext
        ) -> SkillResult:
            return await agent_capability.ReadAttachmentSkillExecutor().execute(
                session, role, request, context
            )

        async def dispatch_text(
            self,
            _session: AsyncSession,
            _role: AgentRole,
            output: str,
            _context: ExecutionContext,
            *,
            exclude: set[str] | None = None,
        ) -> SkillResult:
            captured["output"] = output
            captured["exclude"] = exclude
            return SkillResult(artifacts=[{"file_name": "附件摘要.md"}])

    monkeypatch.setattr(operations, "run_read_attachment", _fake_read)
    monkeypatch.setattr(
        agent_capability.expert_management, "get_expert_execution", _fake_execution
    )
    context = ExecutionContext(
        attachments=(_attach("a.pdf"),), dispatcher=_Dispatcher(), agent_runner=_fake_agent
    )
    res = await agent_capability.execute(
        db, _role(), "【读附件】a.pdf", execution_context=context
    )
    assert captured["exclude"] == {"read_attachment"}
    assert "【交付】" in captured["output"]
    assert len(res.artifacts) == 1 and res.artifacts[0]["file_name"] == "附件摘要.md"


# ── 注册表 ──────────────────────────────────────────────
def test_read_attachment_registered_and_default_on() -> None:
    assert "read_attachment" in REGISTRY
    skill = REGISTRY["read_attachment"]
    assert skill.default_on is True and skill.legacy_executor is not None
    assert skill.executor_factory is not None
    keys = {s.key for s in enabled_skills(_role())}  # tools=[] → 默认全开
    assert "read_attachment" in keys
