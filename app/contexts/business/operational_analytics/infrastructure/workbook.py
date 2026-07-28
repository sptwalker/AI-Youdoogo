"""Versioned Excel workbook parsing owned by Operational Analytics."""

from __future__ import annotations

import unicodedata
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

ColumnType = str


@dataclass(frozen=True)
class ColumnSpec:
    """Map one Excel header to a normalized field and validation rule."""

    header: str
    key: str
    type: ColumnType = "str"
    required: bool = True
    unit: str = ""


@dataclass(frozen=True)
class TemplateSpec:
    """Versioned workbook shape accepted by the ingestion boundary."""

    name: str
    version: str
    columns: tuple[ColumnSpec, ...]
    unique_keys: tuple[str, ...] = ()

    @property
    def headers(self) -> frozenset[str]:
        return frozenset(column.header for column in self.columns)


class TemplateRegistry:
    """Match workbook headers to the most specific registered template."""

    def __init__(self) -> None:
        self._templates: list[TemplateSpec] = []

    def register(self, template: TemplateSpec) -> None:
        self._templates = [
            current
            for current in self._templates
            if (current.name, current.version) != (template.name, template.version)
        ]
        self._templates.append(template)

    def match(self, file_headers: list[str]) -> TemplateSpec | None:
        present = {header.strip() for header in file_headers if header and header.strip()}
        candidates = [template for template in self._templates if template.headers <= present]
        return max(candidates, key=lambda template: len(template.columns)) if candidates else None


registry = TemplateRegistry()
registry.register(
    TemplateSpec(
        name="ops_daily",
        version="v1",
        columns=(
            ColumnSpec("日期", "stat_date", "date"),
            ColumnSpec("产品", "product", "str"),
            ColumnSpec("日活", "dau", "int", unit="人"),
            ColumnSpec("新增", "new_users", "int", required=False, unit="人"),
            ColumnSpec("次留", "retention_d1", "float", required=False, unit="%"),
        ),
        unique_keys=("stat_date", "product"),
    )
)


@dataclass
class RowError:
    row_no: int
    field: str
    message: str


@dataclass
class ParseResult:
    template: str
    rows: list[dict[str, Any]] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    duplicate_count: int = 0

    @property
    def ok_count(self) -> int:
        return len(self.rows)

    @property
    def error_count(self) -> int:
        return len(self.errors)


class ExcelParseError(Exception):
    """Whole-workbook validation failure."""


def _norm_text(value: Any) -> str:
    return unicodedata.normalize("NFKC", str(value)).strip()


def _parse_date(raw: Any, text: str) -> date:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


def _coerce(value: Any, column: ColumnSpec, row_no: int) -> tuple[Any, RowError | None]:
    if value is None or (isinstance(value, str) and not value.strip()):
        if column.required:
            return None, RowError(row_no, column.key, f"必填项「{column.header}」为空")
        return None, None
    text = _norm_text(value)
    try:
        if column.type == "int":
            return int(float(text)), None
        if column.type == "float":
            return float(text), None
        if column.type == "date":
            return _parse_date(value, text), None
        return text, None
    except (TypeError, ValueError):
        return None, RowError(
            row_no,
            column.key,
            f"「{column.header}」值 {text!r} 不是合法 {column.type}",
        )


def _row_is_blank(raw_row: tuple[Any, ...]) -> bool:
    return all(
        value is None or (isinstance(value, str) and not value.strip())
        for value in raw_row
    )


def _build_record(
    raw_row: tuple[Any, ...],
    template: TemplateSpec,
    column_indexes: dict[str, int],
    row_no: int,
) -> tuple[dict[str, Any], list[RowError]]:
    record: dict[str, Any] = {}
    errors: list[RowError] = []
    for column in template.columns:
        index = column_indexes.get(column.key)
        cell = raw_row[index] if index is not None and index < len(raw_row) else None
        value, error = _coerce(cell, column, row_no)
        if error is not None:
            errors.append(error)
        else:
            record[column.key] = value
    return record, errors


def _open_rows(content: bytes) -> tuple[Any, tuple[Any, ...], Iterator[tuple[Any, ...]]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl has no stable failure base class
        raise ExcelParseError(f"无法读取 Excel 文件：{type(exc).__name__}") from exc
    worksheet = workbook.active
    if worksheet is None:
        workbook.close()
        raise ExcelParseError("Excel 无有效工作表")
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        workbook.close()
        raise ExcelParseError("Excel 为空") from None
    return workbook, header_row, rows


def _resolve_columns(
    headers: list[str],
    template: TemplateSpec | None,
) -> tuple[TemplateSpec, dict[str, int]]:
    selected = template or registry.match(headers)
    if selected is None:
        raise ExcelParseError("未匹配到任何已知模板，请使用官方模板上传")
    column_indexes: dict[str, int] = {}
    for column in selected.columns:
        if column.header in headers:
            column_indexes[column.key] = headers.index(column.header)
        elif column.required:
            raise ExcelParseError(f"缺少必填列「{column.header}」")
    return selected, column_indexes


def _is_duplicate(
    record: dict[str, Any],
    template: TemplateSpec,
    seen: set[tuple[object, ...]],
) -> bool:
    if not template.unique_keys:
        return False
    key = tuple(record[field] for field in template.unique_keys)
    if key in seen:
        return True
    seen.add(key)
    return False


def parse_workbook(content: bytes, template: TemplateSpec | None = None) -> ParseResult:
    """Parse an xlsx file, normalizing valid rows and reporting row-level failures."""
    workbook, header_row, rows = _open_rows(content)
    try:
        headers = [_norm_text(value) if value is not None else "" for value in header_row]
        selected, column_indexes = _resolve_columns(headers, template)
        result = ParseResult(template=f"{selected.name}/{selected.version}")
        seen: set[tuple[object, ...]] = set()
        for row_no, raw_row in enumerate(rows, start=2):
            if _row_is_blank(raw_row):
                continue
            record, errors = _build_record(raw_row, selected, column_indexes, row_no)
            if errors:
                result.errors.extend(errors)
                continue
            if _is_duplicate(record, selected, seen):
                result.duplicate_count += 1
                continue
            result.rows.append(record)
        return result
    finally:
        workbook.close()


__all__ = [
    "ColumnSpec",
    "ExcelParseError",
    "ParseResult",
    "RowError",
    "TemplateSpec",
    "parse_workbook",
    "registry",
]
