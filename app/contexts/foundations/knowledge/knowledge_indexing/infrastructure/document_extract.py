"""Document byte extraction adapter owned by Knowledge Indexing.

支持：txt/md（UTF-8 文本）、docx（python-docx）、pdf（pypdf）、粘贴/飞书文档（走 ingest_text）。
"""

from __future__ import annotations

from io import BytesIO

from app.contexts.shared_kernel import ApplicationError, RuleViolation

_TEXT_MIMES = {"text/plain", "text/markdown", "text/x-markdown", ""}
_TEXT_EXTS = (".txt", ".md", ".markdown")


def _extract_docx(content: bytes) -> str:
    """docx → 段落纯文本（含表格单元格文本）。"""
    from docx import Document  # 延迟导入，避免无 docx 上传时也加载 lxml

    doc = Document(BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(content: bytes) -> str:
    """pdf → 逐页抽取纯文本。无文本层（扫描件/图片型 PDF）时明确提示需 OCR。"""
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if not text:
        raise RuleViolation(
            "该 PDF 无文本层（疑似扫描件/图片型），暂不支持；请上传含文本的 PDF 或先 OCR 转文字"
        )
    return text


def extract_text(content: bytes, mime_type: str | None, file_name: str) -> str:
    """把文件字节解成纯文本。支持 txt/md/docx/pdf；其余抛 ApplicationError。"""
    name = file_name.lower()
    mime = (mime_type or "").split(";")[0].strip().lower()
    try:
        if name.endswith(".docx"):
            return _extract_docx(content)
        if name.endswith(".pdf"):
            return _extract_pdf(content)
        if mime in _TEXT_MIMES or name.endswith(_TEXT_EXTS):
            return content.decode("utf-8")
    except ApplicationError:
        raise
    except UnicodeDecodeError as exc:
        raise RuleViolation("文件不是 UTF-8 文本，无法解析") from exc
    except Exception as exc:  # noqa: BLE001 - docx/pdf 解析抛型不稳定，统一成可读错误
        raise RuleViolation(f"文档解析失败：{type(exc).__name__}") from exc
    raise RuleViolation(f"暂不支持的文件类型：{mime_type or file_name}（支持 txt/md/docx/pdf）")
