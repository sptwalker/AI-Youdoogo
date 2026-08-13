"""Pure file rendering and deterministic object-path rules."""

from __future__ import annotations

import csv
import hashlib
import io
import re
import uuid

from openpyxl import Workbook

from app.contexts.foundations.execution.deliverable_management.contracts.delivery import (
    DeliverableFormat,
)

CONTENT_TYPES = {
    DeliverableFormat.CSV: "text/csv; charset=utf-8",
    DeliverableFormat.XLSX: (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ),
    DeliverableFormat.MARKDOWN: "text/markdown; charset=utf-8",
    DeliverableFormat.TEXT: "text/plain; charset=utf-8",
    DeliverableFormat.PPTX: (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    ),
}


def markdown_table_to_rows(body: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for raw_line in body.splitlines():
        line = raw_line.strip()
        if not line or "|" not in line:
            continue
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(set(cell) <= {"-", ":", " "} and cell for cell in cells):
            continue
        rows.append(cells)
    return rows


def markdown_to_slides(body: str) -> list[tuple[str, list[str]]]:
    """把 markdown 正文切成幻灯片：#/##/… 行起新页作标题，其余非空行去掉项目符号作要点。

    正文首行若非标题则自成首页标题（保证任意正文都能出至少一页）。纯确定性，无 LLM。
    # ponytail: 表格数据行仍逐行作纯文本要点（保留「| 指标 | 值 |」原样，仅剔除 --- 分隔行）；
    #   要真渲染成 pptx 表格再引 python-pptx 的 add_table，当前按文本够用、避免过度实现。
    """
    slides: list[tuple[str, list[str]]] = []
    for raw_line in body.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if "|" in line and all(
            set(cell) <= {"-", ":", " "} and cell for cell in line.strip("|").split("|")
        ):
            continue  # 跳过 Markdown 表格分隔行「| --- | :--: |」，否则会变成一条空要点
        heading = re.match(r"^#{1,6}\s+(.*)$", line)
        if heading:
            slides.append((heading.group(1).strip(), []))
        elif slides:
            slides[-1][1].append(re.sub(r"^[-*•]\s*", "", line))
        else:
            slides.append((line, []))
    return slides


def build_bytes(file_format: DeliverableFormat | str, body: str) -> bytes:
    normalized = DeliverableFormat(file_format)
    if normalized in (DeliverableFormat.MARKDOWN, DeliverableFormat.TEXT):
        return body.strip().encode("utf-8")
    if normalized is DeliverableFormat.PPTX:
        slides = markdown_to_slides(body)
        if not slides:
            raise ValueError("交付内容为空，无法生成 PPT")
        from pptx import Presentation

        presentation = Presentation()
        layout = presentation.slide_layouts[1]  # 标题 + 内容 版式
        for title, bullets in slides:
            slide = presentation.slides.add_slide(layout)
            slide.shapes.title.text = title
            frame = slide.placeholders[1].text_frame
            frame.text = bullets[0] if bullets else ""
            for bullet in bullets[1:]:
                frame.add_paragraph().text = bullet
        output = io.BytesIO()
        presentation.save(output)
        return output.getvalue()
    rows = markdown_table_to_rows(body)
    if not rows:
        raise ValueError("交付内容不含可解析的表格")
    if normalized is DeliverableFormat.CSV:
        buffer = io.StringIO()
        csv.writer(buffer).writerows(rows)
        return buffer.getvalue().encode("utf-8-sig")
    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def safe_file_name(name: str, file_format: DeliverableFormat | str) -> str:
    normalized = DeliverableFormat(file_format)
    base = re.sub(r"[/\\:*?\"<>|]", "_", name).strip()[:120] or "交付物"
    suffix = f".{normalized.value}"
    return base if base.lower().endswith(suffix) else base + suffix


def object_name(
    deliverable_id: uuid.UUID,
    file_name: str,
    idempotency_key: str | None,
) -> str:
    if idempotency_key:
        digest = hashlib.sha256(idempotency_key.encode()).hexdigest()
        return f"deliverables/idempotent/{digest}/{file_name}"
    return f"deliverables/{deliverable_id}/{file_name}"
