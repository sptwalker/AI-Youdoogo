"""附件字节 → 纯文本解析（按扩展名分派 pdf/docx/xlsx/文本）。

只做「字节进、文本出」，不碰 MinIO/DB（取字节在 operations 层），故可离线单测。
扫描件 PDF（无文本层）→ 返「疑似扫描件」原因；不支持的格式返原因；解析异常返空 + 原因、不抛穿
（仿 read_url，附件解析失败不打断桌面消息流）。文本截断上限防超大文档灌爆模型。
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass

logger = logging.getLogger(__name__)

_MAX_TEXT = 12000  # 抽取正文截断（喂回模型的素材上限，附件通常比网页长故略宽）
_TEXT_EXTENSIONS = (".txt", ".md", ".csv", ".log", ".json")


@dataclass(frozen=True, slots=True)
class AttachmentOutcome:
    """归一化正文 + 结构化失败原因（交由上层折进 note，不抛穿）。"""

    text: str
    reason: str = ""


def _ext(name: str) -> str:
    _, _, tail = name.lower().rpartition(".")
    return f".{tail}" if tail and "." in name else ""


def _clip(text: str) -> str:
    stripped = text.strip()
    return stripped[:_MAX_TEXT]


def _parse_pdf(data: bytes) -> AttachmentOutcome:
    from pypdf import PdfReader  # 局部导入：缺依赖时只影响该格式，不拖垮模块加载

    reader = PdfReader(io.BytesIO(data))
    parts = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(p for p in parts if p.strip())
    if not text.strip():
        # 扫描件/图片型 PDF 无文本层 → extract_text 全空
        return AttachmentOutcome(text="", reason="疑似扫描件或图片型 PDF，无可提取文本")
    return AttachmentOutcome(text=_clip(text))


def _parse_docx(data: bytes) -> AttachmentOutcome:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
    if not text.strip():
        return AttachmentOutcome(text="", reason="文档为空或无可提取文本")
    return AttachmentOutcome(text=_clip(text))


def _parse_xlsx(data: bytes) -> AttachmentOutcome:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"# 工作表：{sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                lines.append("\t".join(cells))
    workbook.close()
    text = "\n".join(lines)
    if not text.strip():
        return AttachmentOutcome(text="", reason="表格为空或无可提取内容")
    return AttachmentOutcome(text=_clip(text))


def _parse_text(data: bytes) -> AttachmentOutcome:
    text = data.decode("utf-8", errors="replace")
    if not text.strip():
        return AttachmentOutcome(text="", reason="文件为空")
    return AttachmentOutcome(text=_clip(text))


def parse_attachment(name: str, data: bytes) -> AttachmentOutcome:
    """按文件名扩展名分派解析；不支持的格式与异常都返空 + 原因、不抛穿。"""
    ext = _ext(name)
    try:
        if ext == ".pdf":
            return _parse_pdf(data)
        if ext == ".docx":
            return _parse_docx(data)
        if ext == ".xlsx":
            return _parse_xlsx(data)
        if ext in _TEXT_EXTENSIONS:
            return _parse_text(data)
    except Exception as exc:  # noqa: BLE001 - 解析失败折进原因，不打断消息流
        logger.warning("附件解析失败 ext=%s err=%s", ext, type(exc).__name__)
        return AttachmentOutcome(text="", reason=f"附件解析异常（{type(exc).__name__}）")
    return AttachmentOutcome(text="", reason=f"暂不支持该格式（{ext or '无扩展名'}）")
