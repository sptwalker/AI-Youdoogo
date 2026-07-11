"""Excel 解析管道：识别模板 → 逐行校验/标准化 → 去重 → 汇总结果（docs/11 A类）。

纯逻辑，无数据库依赖：入库由上层 service 消费 ParseResult 完成。
治理动作：类型转换、必填校验、前后空白与全半角归一、批内按业务唯一键去重。
"""

from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services.excel_ingest.spec import ColumnSpec, TemplateSpec, registry


@dataclass
class RowError:
    """一行的校验失败记录（row_no 为 Excel 行号，含表头从 1 计）。"""

    row_no: int
    field: str
    message: str


@dataclass
class ParseResult:
    """解析结果汇总。"""

    template: str
    rows: list[dict[str, Any]] = field(default_factory=list)  # 校验通过、已标准化、已去重
    errors: list[RowError] = field(default_factory=list)
    duplicate_count: int = 0

    @property
    def ok_count(self) -> int:
        return len(self.rows)

    @property
    def error_count(self) -> int:
        return len(self.errors)


class ExcelParseError(Exception):
    """整表级失败（文件损坏、无法匹配模板等），非单行错误。"""


def _norm_text(value: Any) -> str:
    """全半角归一 + 去首尾空白（NFKC 把全角数字/字母/空格折成半角）。"""
    return unicodedata.normalize("NFKC", str(value)).strip()


def _coerce(value: Any, col: ColumnSpec, row_no: int) -> tuple[Any, RowError | None]:
    """按列类型转换一个单元格；返回 (值, 错误)。空值在必填时报错。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        if col.required:
            return None, RowError(row_no, col.key, f"必填项「{col.header}」为空")
        return None, None

    text = _norm_text(value)
    try:
        if col.type == "int":
            return int(float(text)), None  # "12.0" → 12
        if col.type == "float":
            return float(text), None
        if col.type == "date":
            return _parse_date(value, text), None
        return text, None
    except (ValueError, TypeError):
        return None, RowError(row_no, col.key, f"「{col.header}」值 {text!r} 不是合法 {col.type}")


def _parse_date(raw: Any, text: str) -> date:
    """接受 datetime/date（openpyxl 原生）或 YYYY-MM-DD / YYYY/MM/DD 文本。"""
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


def parse_workbook(content: bytes, template: TemplateSpec | None = None) -> ParseResult:
    """解析 xlsx 字节流。template 为空时按首行表头自动匹配注册表。

    Raises:
        ExcelParseError: 文件无法读取、为空、或无法匹配到模板。
    """
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl 抛型不稳定，统一成整表错误
        raise ExcelParseError(f"无法读取 Excel 文件：{type(exc).__name__}") from exc

    ws = wb.active
    if ws is None:
        raise ExcelParseError("Excel 无有效工作表")

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ExcelParseError("Excel 为空") from None

    headers = [_norm_text(h) if h is not None else "" for h in header_row]
    tpl = template or registry.match(headers)
    if tpl is None:
        raise ExcelParseError("未匹配到任何已知模板，请使用官方模板上传")

    # 表头文本 → 列索引
    col_index: dict[str, int] = {}
    for col in tpl.columns:
        if col.header in headers:
            col_index[col.key] = headers.index(col.header)
        elif col.required:
            raise ExcelParseError(f"缺少必填列「{col.header}」")

    result = ParseResult(template=f"{tpl.name}/{tpl.version}")
    seen: set[tuple] = set()

    for offset, raw_row in enumerate(rows_iter, start=2):  # 数据从第2行起
        if _row_is_blank(raw_row):
            continue
        record, row_errors = _build_record(raw_row, tpl, col_index, offset)
        if row_errors:
            result.errors.extend(row_errors)
            continue
        if tpl.unique_keys:
            key = tuple(record[k] for k in tpl.unique_keys)
            if key in seen:
                result.duplicate_count += 1
                continue
            seen.add(key)
        result.rows.append(record)

    return result


def _row_is_blank(raw_row: tuple) -> bool:
    return all(c is None or (isinstance(c, str) and not c.strip()) for c in raw_row)


def _build_record(
    raw_row: tuple, tpl: TemplateSpec, col_index: dict[str, int], row_no: int
) -> tuple[dict[str, Any], list[RowError]]:
    record: dict[str, Any] = {}
    errors: list[RowError] = []
    for col in tpl.columns:
        idx = col_index.get(col.key)
        cell = raw_row[idx] if idx is not None and idx < len(raw_row) else None
        value, err = _coerce(cell, col, row_no)
        if err:
            errors.append(err)
        else:
            record[col.key] = value
    return record, errors
